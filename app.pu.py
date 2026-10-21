"""
Member Database & Monthly Dues Management System — single-file version.
Run with:  streamlit run app.py
"""
from __future__ import annotations

import types

import io
import json
import sqlite3
import datetime as dt
from pathlib import Path

import pandas as pd

APP_DIR = Path(__file__).parent
DB_PATH = APP_DIR / "members.db"
UNI_FILE = APP_DIR / "Chongqing_Universities_Locations.xlsx"

BUILTIN_UNIVERSITIES = [
    ('Chongqing University', '重庆大学', 'Public', 'Shapingba District (沙坪坝区)', '174 Shazheng St, Shapingba District', '400044'),
    ('Southwest University', '西南大学', 'Public', 'Beibei District (北碚区)', '2 Tiansheng Rd, Beibei District', '400716'),
    ('Chongqing University of Posts and Telecommunications', '重庆邮电大学', 'Public', "Nan'an District (南岸区)", "2 Chongwen Rd, Nan'an District", '400065'),
    ('Chongqing Jiaotong University', '重庆交通大学', 'Public', "Nan'an District (南岸区)", "66 Xuefu Ave, Nan'an District", '400074'),
    ('Chongqing Medical University', '重庆医科大学', 'Public', 'Yuzhong District (渝中区)', '1 Yixueyuan Rd, Yuzhong District', '400016'),
    ('Chongqing Normal University', '重庆师范大学', 'Public', 'Shapingba District (沙坪坝区)', '37 Daxuecheng Middle Rd, Shapingba District', '401331'),
    ('Chongqing University of Arts and Sciences', '重庆文理学院', 'Public', 'Yongchuan District (永川区)', '319 Honghe Ave, Yongchuan District', '402160'),
    ('Chongqing Three Gorges University', '重庆三峡学院', 'Public', 'Wanzhou District (万州区)', '666 Tianxing Rd, Wanzhou District', '404000'),
    ('Yangtze Normal University', '长江师范学院', 'Public', 'Fuling District (涪陵区)', '16 Juxian Ave, Fuling District', '408100'),
    ('Sichuan International Studies University', '四川外国语大学', 'Public', 'Shapingba District (沙坪坝区)', '33 Zhuangzhi Rd, Shapingba District', '400031'),
    ('Southwest University of Political Science & Law', '西南政法大学', 'Public', 'Yubei District (渝北区)', '301 Baosheng Ave, Yubei District', '401120'),
    ('Sichuan Fine Arts Institute', '四川美术学院', 'Public', 'Shapingba District (沙坪坝区)', '56 Huxi University Town South Rd, Shapingba District', '401331'),
    ('Chongqing University of Science and Technology', '重庆科技大学', 'Public', 'Shapingba District (沙坪坝区)', '20 Daxuecheng East Rd, Shapingba District', '401331'),
    ('Chongqing University of Technology', '重庆理工大学', 'Public', 'Banan District (巴南区)', '69 Hongguang Ave, Banan District', '400054'),
    ('Chongqing Technology and Business University', '重庆工商大学', 'Public', "Nan'an District (南岸区)", "19 Xuefu Ave, Nan'an District", '400067'),
    ('Chongqing Police College', '重庆警察学院', 'Public', 'Shapingba District (沙坪坝区)', '666 Jingzheng Rd, Shapingba District', '401331'),
    ('Chongqing University of Education', '重庆第二师范学院', 'Public', "Nan'an District (南岸区)", "1 Chongjiao Rd, Nanshan Subdistrict, Nan'an District", '400067'),
    ('Chongqing University of Chinese Medicine', '重庆中医药学院', 'Public', 'Bishan District (璧山区)', '61 Puguobao Rd, Bishan District', '402760'),
    ('Chongqing Electronic Technology Vocational University', '重庆电子科技职业大学', 'Public', 'Shapingba District (沙坪坝区)', '76 Daxuecheng East Rd, Shapingba District', '401331'),
    ('Chongqing Industrial Polytechnic University', '重庆工业职业技术大学', 'Public', 'Yubei District (渝北区)', '1000 Taoyuan Ave, Konggang, Yubei District', '401120'),
    ('Chongqing Vocational University of Mechanical and Electrical Technology', '重庆机电职业技术大学', 'Private', 'Bishan District (璧山区)', '1001 Biqing North Rd, Bishan District', '402760'),
    ('Chongqing Institute of Engineering', '重庆工程学院', 'Private', 'Banan District (巴南区)', '16 Baihelin, Nanquan Subdistrict, Banan District', '400056'),
    ('Chongqing Metropolitan College of Science and Technology', '重庆城市科技学院', 'Private', 'Yongchuan District (永川区)', '368 Guangcai Ave, Yongchuan District', '402167'),
    ('Chongqing College of Humanities, Science & Technology', '重庆人文科技学院', 'Private', 'Hechuan District (合川区)', '256 Xueyuan St, Caojie Subdistrict, Hechuan District', '401524'),
    ('Chongqing Institute of Foreign Studies', '重庆外语外事学院', 'Private', 'Yubei District (渝北区)', '18 Longshi Rd, Yubei District', '401120'),
    ('Chongqing College of International Business and Economics', '重庆对外经贸学院', 'Private', 'Hechuan District (合川区)', '9 Xuefu Rd, Hechuan District', '401520'),
    ('Chongqing College of Finance and Economics', '重庆财经学院', 'Private', 'Banan District (巴南区)', '906 Shangwen Ave, Longzhou Bay, Banan District', '401320'),
    ('Pass College of Chongqing Technology and Business University', '重庆工商大学派斯学院', 'Private', 'Hechuan District (合川区)', '593 Jiaotong St, Hechuan District', '401520'),
    ('Chongqing University of Mobile Communication', '重庆移通学院', 'Private', 'Hechuan District (合川区)', '1 Holiday Ave, University Town, Hechuan District', '401520'),
]

# --------------------------------------------------------------------------------------
# Fixed choice lists used in dropdowns
# --------------------------------------------------------------------------------------
GENDERS = ["Male", "Female", "Prefer not to say"]
DEGREE_LEVELS = ["Language Program (Chinese)", "Diploma / Vocational", "Bachelor's",
                 "Master's", "PhD", "Post-Doctoral", "Exchange / Visiting", "Staff / Faculty", "Other"]
SCHOLARSHIPS = ["Chinese Government Scholarship (CSC)", "Provincial / City Scholarship",
                "University Scholarship", "Confucius Institute Scholarship",
                "Home Government Scholarship", "Self-Sponsored", "Other"]
RESIDENCE = ["On-campus dormitory", "Off-campus apartment", "Other"]
STATUSES = ["Active", "Inactive", "Graduated", "Suspended", "Pending approval"]
PENDING = "Pending approval"
EXIT_STATUSES = ["Inactive", "Graduated", "Suspended"]
PAY_METHODS = ["Cash", "WeChat Pay", "Alipay", "Bank Transfer", "Mobile Money", "Other"]

DEFAULT_SETTINGS = {
    "org_name": "Members Association – Chongqing",
    "currency": "¥",
    "membership_types": json.dumps({"Regular": 50, "Executive": 50, "Associate": 30, "Honorary": 0}),
    "due_day": "10",  # day of month dues are expected by
    "app_url": "",
}

MEMBER_FIELDS = [
    "full_name", "gender", "date_of_birth", "nationality", "phone", "email", "wechat_id",
    "university_id", "student_id", "degree_level", "field_of_study", "enrollment_year",
    "expected_graduation", "scholarship", "residence", "membership_type", "monthly_due",
    "join_date", "status", "exit_date", "emergency_name", "emergency_phone", "notes",
]

SCHEMA = """
CREATE TABLE IF NOT EXISTS universities(
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    name_en TEXT UNIQUE NOT NULL,
    name_cn TEXT, ownership TEXT, district TEXT, address TEXT, postal_code TEXT
);
CREATE TABLE IF NOT EXISTS members(
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    member_code TEXT UNIQUE,
    full_name TEXT NOT NULL, gender TEXT, date_of_birth TEXT, nationality TEXT,
    phone TEXT, email TEXT, wechat_id TEXT,
    university_id INTEGER REFERENCES universities(id),
    student_id TEXT, degree_level TEXT, field_of_study TEXT,
    enrollment_year INTEGER, expected_graduation INTEGER,
    scholarship TEXT, residence TEXT,
    membership_type TEXT, monthly_due REAL NOT NULL DEFAULT 0,
    join_date TEXT NOT NULL, status TEXT DEFAULT 'Active', exit_date TEXT,
    emergency_name TEXT, emergency_phone TEXT, notes TEXT,
    created_at TEXT DEFAULT CURRENT_TIMESTAMP, updated_at TEXT
);
CREATE TABLE IF NOT EXISTS payments(
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    member_id INTEGER NOT NULL REFERENCES members(id) ON DELETE CASCADE,
    period TEXT NOT NULL,             -- 'YYYY-MM' the payment covers
    amount REAL NOT NULL,
    payment_date TEXT NOT NULL,
    method TEXT, reference TEXT, received_by TEXT, notes TEXT,
    created_at TEXT DEFAULT CURRENT_TIMESTAMP
);
CREATE INDEX IF NOT EXISTS ix_pay_member ON payments(member_id);
CREATE INDEX IF NOT EXISTS ix_pay_period ON payments(period);
CREATE TABLE IF NOT EXISTS settings(key TEXT PRIMARY KEY, value TEXT);
"""

MEMBER_COLS = ["id", "member_code", *MEMBER_FIELDS, "created_at", "updated_at"]
PAYMENT_COLS = ["id", "member_id", "period", "amount", "payment_date", "method",
                "reference", "received_by", "notes", "created_at"]
UNI_COLS = ["id", "name_en", "name_cn", "ownership", "district", "address", "postal_code"]


# --------------------------------------------------------------------------------------
# Connection / setup
# --------------------------------------------------------------------------------------
def connect(path: Path | str | None = None) -> sqlite3.Connection:
    conn = sqlite3.connect(str(path or DB_PATH), check_same_thread=False)
    conn.execute("PRAGMA foreign_keys = ON")
    conn.executescript(SCHEMA)
    for k, v in DEFAULT_SETTINGS.items():
        conn.execute("INSERT OR IGNORE INTO settings(key, value) VALUES (?, ?)", (k, v))
    conn.commit()
    if conn.execute("SELECT COUNT(*) FROM universities").fetchone()[0] == 0:
        if UNI_FILE.exists():
            import_universities(conn, UNI_FILE)
        else:  # built-in copy of the Chongqing universities list
            conn.executemany("INSERT OR IGNORE INTO universities(name_en,name_cn,ownership,district,address,"
                             "postal_code) VALUES (?,?,?,?,?,?)", BUILTIN_UNIVERSITIES)
            conn.commit()
    return conn


def _clean(v):
    if v is None:
        return None
    if isinstance(v, float) and pd.isna(v):
        return None
    s = str(v).strip()
    if s.endswith(".0") and s[:-2].isdigit():  # postal codes read as floats
        s = s[:-2]
    return s or None


def import_universities(conn, file) -> int:
    """Load / refresh the university list from the Excel file. Returns rows upserted."""
    df = pd.read_excel(file)
    df.columns = [str(c).strip().lower() for c in df.columns]

    def pick(*keys):
        for c in df.columns:
            if any(k in c for k in keys):
                return c
        return None

    c_en = pick("english", "university")
    c_cn = pick("chinese")
    c_own = pick("ownership", "type")
    c_dist = pick("district", "area")
    c_addr = pick("address")
    c_post = pick("postal", "zip")
    if c_en is None:
        raise ValueError("Could not find a university name column in the file.")
    n = 0
    for _, r in df.iterrows():
        name = _clean(r[c_en])
        if not name:
            continue
        vals = (name,
                _clean(r[c_cn]) if c_cn else None,
                _clean(r[c_own]) if c_own else None,
                _clean(r[c_dist]) if c_dist else None,
                _clean(r[c_addr]) if c_addr else None,
                _clean(r[c_post]) if c_post else None)
        conn.execute(
            """INSERT INTO universities(name_en, name_cn, ownership, district, address, postal_code)
               VALUES (?,?,?,?,?,?)
               ON CONFLICT(name_en) DO UPDATE SET name_cn=excluded.name_cn, ownership=excluded.ownership,
               district=excluded.district, address=excluded.address, postal_code=excluded.postal_code""",
            vals)
        n += 1
    conn.commit()
    return n


# --------------------------------------------------------------------------------------
# Settings
# --------------------------------------------------------------------------------------
def get_settings(conn) -> dict:
    s = dict(conn.execute("SELECT key, value FROM settings").fetchall())
    s["membership_types"] = json.loads(s.get("membership_types") or "{}")
    return s


def save_setting(conn, key, value):
    if isinstance(value, (dict, list)):
        value = json.dumps(value)
    conn.execute("INSERT INTO settings(key,value) VALUES(?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",
                 (key, str(value)))
    conn.commit()


# --------------------------------------------------------------------------------------
# Universities
# --------------------------------------------------------------------------------------
def universities(conn) -> pd.DataFrame:
    return pd.read_sql("SELECT * FROM universities ORDER BY name_en", conn)


def add_university(conn, name_en, name_cn, ownership, district, address, postal_code):
    conn.execute("INSERT INTO universities(name_en,name_cn,ownership,district,address,postal_code) VALUES(?,?,?,?,?,?)",
                 (name_en.strip(), name_cn, ownership, district, address, postal_code))
    conn.commit()


# --------------------------------------------------------------------------------------
# Members
# --------------------------------------------------------------------------------------
def _next_code(conn) -> str:
    row = conn.execute("SELECT MAX(id) FROM members").fetchone()[0] or 0
    while True:
        row += 1
        code = f"MBR-{row:04d}"
        if not conn.execute("SELECT 1 FROM members WHERE member_code=?", (code,)).fetchone():
            return code


def _norm(data: dict) -> dict:
    out = {}
    for k in MEMBER_FIELDS:
        v = data.get(k)
        if isinstance(v, (dt.date, dt.datetime)):
            v = v.isoformat()[:10]
        if isinstance(v, str):
            v = v.strip() or None
        out[k] = v
    return out


def validate_member(conn, data: dict, member_id: int | None = None) -> list[str]:
    errs = []
    if not (data.get("full_name") or "").strip():
        errs.append("Full name is required.")
    if not data.get("university_id"):
        errs.append("Please select a university.")
    if not data.get("join_date"):
        errs.append("Join date is required.")
    email = (data.get("email") or "").strip()
    if email and ("@" not in email or "." not in email.split("@")[-1]):
        errs.append("Email address does not look valid.")
    if (data.get("monthly_due") or 0) < 0:
        errs.append("Monthly due cannot be negative.")
    if data.get("status") in EXIT_STATUSES and not data.get("exit_date"):
        errs.append("Please give the date the member left / graduated (billing stops then).")
    if data.get("exit_date") and data.get("join_date") and str(data["exit_date"]) < str(data["join_date"]):
        errs.append("Exit date cannot be before join date.")
    sid = (data.get("student_id") or "").strip()
    if sid:
        q = "SELECT member_code FROM members WHERE student_id=? AND university_id=?"
        args = [sid, data.get("university_id")]
        if member_id:
            q += " AND id<>?"
            args.append(member_id)
        hit = conn.execute(q, args).fetchone()
        if hit:
            errs.append(f"Student ID {sid} is already registered at this university ({hit[0]}).")
    return errs


def add_member(conn, data: dict) -> str:
    d = _norm(data)
    code = _next_code(conn)
    cols = ["member_code", *MEMBER_FIELDS]
    conn.execute(f"INSERT INTO members({','.join(cols)}) VALUES ({','.join('?' * len(cols))})",
                 [code, *[d[k] for k in MEMBER_FIELDS]])
    conn.commit()
    return code


def update_member(conn, member_id: int, data: dict):
    d = _norm(data)
    sets = ", ".join(f"{k}=?" for k in MEMBER_FIELDS)
    conn.execute(f"UPDATE members SET {sets}, updated_at=CURRENT_TIMESTAMP WHERE id=?",
                 [*[d[k] for k in MEMBER_FIELDS], member_id])
    conn.commit()


def delete_member(conn, member_id: int):
    conn.execute("DELETE FROM payments WHERE member_id=?", (member_id,))
    conn.execute("DELETE FROM members WHERE id=?", (member_id,))
    conn.commit()


def members(conn) -> pd.DataFrame:
    """Members joined with their university details."""
    df = pd.read_sql(
        """SELECT m.*, u.name_en AS university, u.name_cn AS university_cn, u.ownership,
                  u.district, u.address AS university_address, u.postal_code
           FROM members m LEFT JOIN universities u ON u.id = m.university_id
           ORDER BY m.id""", conn)
    if df.empty:
        df = pd.DataFrame(columns=MEMBER_COLS + ["university", "university_cn", "ownership", "district",
                                                 "university_address", "postal_code"])
    return df


def get_member(conn, member_id: int) -> dict | None:
    df = members(conn)
    row = df[df["id"] == member_id]
    return None if row.empty else row.iloc[0].to_dict()


def approve_member(conn, member_id: int, join_date=None):
    join_date = (join_date or dt.date.today())
    if isinstance(join_date, (dt.date, dt.datetime)):
        join_date = join_date.isoformat()[:10]
    conn.execute("UPDATE members SET status='Active', join_date=?, updated_at=CURRENT_TIMESTAMP WHERE id=?",
                 (join_date, member_id))
    conn.commit()


def portal_lookup(conn, member_code: str, phone_last4: str) -> dict | None:
    """Self-service lookup: member code + last 4 digits of the registered phone number."""
    code = (member_code or "").strip().upper()
    last4 = "".join(ch for ch in (phone_last4 or "") if ch.isdigit())
    if not code or len(last4) != 4:
        return None
    df = members(conn)
    hit = df[df["member_code"].str.upper() == code]
    if hit.empty:
        return None
    m = hit.iloc[0].to_dict()
    digits = "".join(ch for ch in str(m.get("phone") or "") if ch.isdigit())
    return m if len(digits) >= 4 and digits[-4:] == last4 else None


# --------------------------------------------------------------------------------------
# QR code
# --------------------------------------------------------------------------------------
def member_link(base_url: str) -> str:
    base = (base_url or "").strip().rstrip("/")
    if base and not base.startswith(("http://", "https://")):
        base = "https://" + base
    return f"{base}/?page=member" if base else ""


def qr_image(data: str, color: str = "#0D2B55"):
    import qrcode
    from qrcode.constants import ERROR_CORRECT_M
    qr = qrcode.QRCode(error_correction=ERROR_CORRECT_M, box_size=14, border=3)
    qr.add_data(data)
    qr.make(fit=True)
    return qr.make_image(fill_color=color, back_color="white").convert("RGB")


def _font(size: int, bold=False):
    from PIL import ImageFont
    names = (["DejaVuSans-Bold.ttf", "Arial Bold.ttf", "arialbd.ttf"] if bold
             else ["DejaVuSans.ttf", "Arial.ttf", "arial.ttf"])
    for n in names:
        try:
            return ImageFont.truetype(n, size)
        except OSError:
            pass
    try:
        return ImageFont.load_default(size=size)
    except TypeError:
        return ImageFont.load_default()


def qr_poster(qr_img, org_name: str, url: str, headline="Scan to join or check your dues") -> bytes:
    """A4-proportioned printable poster (PNG) with the QR code in the middle."""
    from PIL import Image, ImageDraw
    W, H = 1240, 1754
    navy, gold = (13, 43, 85), (201, 168, 76)
    img = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(img)
    d.rectangle([0, 0, W, 300], fill=navy)
    d.rectangle([0, 300, W, 316], fill=gold)

    def center(text, y, font, fill):
        w = d.textlength(text, font=font)
        d.text(((W - w) / 2, y), text, font=font, fill=fill)

    def fit(text, size, bold, maxw):
        f = _font(size, bold)
        while d.textlength(text, font=f) > maxw and size > 18:
            size -= 2
            f = _font(size, bold)
        return f

    center(org_name, 95, fit(org_name, 64, True, W - 120), "white")
    center("Member Portal", 190, _font(40), gold)
    center(headline, 400, fit(headline, 56, True, W - 120), navy)
    q = qr_img.resize((820, 820), Image.NEAREST)
    x, y = (W - 820) // 2, 520
    d.rounded_rectangle([x - 30, y - 30, x + 850, y + 850], radius=30, outline=gold, width=6)
    img.paste(q, (x, y))
    center("1. Open your phone camera or WeChat scanner", 1430, _font(36), (55, 65, 81))
    center("2. Point it at the code and tap the link", 1485, _font(36), (55, 65, 81))
    center(url, 1600, fit(url, 30, False, W - 120), (107, 114, 128))
    buf = io.BytesIO()
    img.save(buf, format="PNG", dpi=(150, 150))
    return buf.getvalue()


def png_bytes(img) -> bytes:
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


# --------------------------------------------------------------------------------------
# Payments
# --------------------------------------------------------------------------------------
def add_payments(conn, member_id: int, periods: list[str], amount_each: float, payment_date,
                 method=None, reference=None, received_by=None, notes=None) -> int:
    if isinstance(payment_date, (dt.date, dt.datetime)):
        payment_date = payment_date.isoformat()[:10]
    for p in periods:
        conn.execute("""INSERT INTO payments(member_id, period, amount, payment_date, method, reference,
                        received_by, notes) VALUES (?,?,?,?,?,?,?,?)""",
                     (member_id, p, float(amount_each), payment_date, method, reference, received_by, notes))
    conn.commit()
    return len(periods)


def delete_payment(conn, payment_id: int):
    conn.execute("DELETE FROM payments WHERE id=?", (payment_id,))
    conn.commit()


def payments(conn) -> pd.DataFrame:
    df = pd.read_sql(
        """SELECT p.*, m.member_code, m.full_name, u.name_en AS university
           FROM payments p JOIN members m ON m.id = p.member_id
           LEFT JOIN universities u ON u.id = m.university_id
           ORDER BY p.payment_date DESC, p.id DESC""", conn)
    if df.empty:
        df = pd.DataFrame(columns=PAYMENT_COLS + ["member_code", "full_name", "university"])
    df["amount"] = pd.to_numeric(df["amount"], errors="coerce").fillna(0.0)
    return df


# --------------------------------------------------------------------------------------
# Dues engine
# --------------------------------------------------------------------------------------
def period_of(d) -> str:
    if isinstance(d, str):
        return d[:7]
    return f"{d.year:04d}-{d.month:02d}"


def period_range(start: str, end: str) -> list[str]:
    """Inclusive list of 'YYYY-MM' from start to end."""
    if not start or not end or start > end:
        return []
    y, m = int(start[:4]), int(start[5:7])
    ey, em = int(end[:4]), int(end[5:7])
    out = []
    while (y, m) <= (ey, em):
        out.append(f"{y:04d}-{m:02d}")
        m += 1
        if m == 13:
            y, m = y + 1, 1
    return out


def shift_period(p: str, n: int) -> str:
    y, m = int(p[:4]), int(p[5:7])
    idx = y * 12 + (m - 1) + n
    return f"{idx // 12:04d}-{idx % 12 + 1:02d}"


def period_label(p: str) -> str:
    return dt.date(int(p[:4]), int(p[5:7]), 1).strftime("%b %Y")


def billable_periods(member: dict, upto: str) -> list[str]:
    """Months a member owes dues for, from joining month to `upto` (or exit month if earlier)."""
    if member.get("status") == PENDING:  # not billed until an admin approves them
        return []
    start = period_of(member["join_date"]) if member.get("join_date") else None
    end = upto
    exit_d = member.get("exit_date")
    if exit_d and isinstance(exit_d, str) and exit_d.strip():
        end = min(end, period_of(exit_d))
    return period_range(start, end) if start else []


def member_ledger(member: dict, pay_df: pd.DataFrame, upto: str) -> pd.DataFrame:
    """Month-by-month statement for one member."""
    due = float(member.get("monthly_due") or 0)
    periods = billable_periods(member, upto)
    mp = pay_df[pay_df["member_id"] == member["id"]]
    paid_by = mp.groupby("period")["amount"].sum().to_dict()
    # include prepaid / extra months that fall outside the billable window
    extra = sorted(p for p in paid_by if p not in periods)
    rows, running = [], 0.0
    for p in sorted(set(periods) | set(extra)):
        d = due if p in periods else 0.0
        paid = float(paid_by.get(p, 0.0))
        bal = d - paid
        running += bal
        if d == 0 and paid > 0:
            status = "Advance / Extra"
        elif d == 0:
            status = "No due"
        elif paid >= d:
            status = "Paid"
        elif paid > 0:
            status = "Partial"
        else:
            status = "Unpaid"
        rows.append({"Period": p, "Month": period_label(p), "Due": d, "Paid": paid,
                     "Balance": bal, "Running Balance": running, "Status": status})
    return pd.DataFrame(rows, columns=["Period", "Month", "Due", "Paid", "Balance", "Running Balance", "Status"])


def dues_summary(conn, upto: str, mem_df: pd.DataFrame | None = None,
                 pay_df: pd.DataFrame | None = None) -> pd.DataFrame:
    """One row per member: expected vs paid up to `upto`, balance, months in arrears."""
    mem_df = members(conn) if mem_df is None else mem_df
    pay_df = payments(conn) if pay_df is None else pay_df
    cols = ["id", "member_code", "full_name", "university", "district", "phone", "email",
            "membership_type", "status", "monthly_due", "months_billed", "expected", "paid",
            "balance", "months_in_arrears", "oldest_unpaid", "last_payment"]
    rows = []
    for m in mem_df.to_dict("records"):
        led = member_ledger(m, pay_df, upto)
        led_upto = led[led["Period"] <= upto]
        mp = pay_df[pay_df["member_id"] == m["id"]]
        expected = float(led_upto["Due"].sum())
        paid = float(mp["amount"].sum())  # all payments count as credit, incl. advance
        unpaid = led_upto[(led_upto["Due"] > 0) & (led_upto["Paid"] < led_upto["Due"])]
        rows.append({
            "id": m["id"], "member_code": m["member_code"], "full_name": m["full_name"],
            "university": m.get("university"), "district": m.get("district"),
            "phone": m.get("phone"), "email": m.get("email"),
            "membership_type": m.get("membership_type"), "status": m.get("status"),
            "monthly_due": float(m.get("monthly_due") or 0),
            "months_billed": int((led_upto["Due"] > 0).sum()),
            "expected": expected, "paid": paid, "balance": expected - paid,
            "months_in_arrears": len(unpaid),
            "oldest_unpaid": period_label(unpaid["Period"].min()) if len(unpaid) else "",
            "last_payment": mp["payment_date"].max() if len(mp) else "",
        })
    return pd.DataFrame(rows, columns=cols)


def monthly_collection(conn, period: str, mem_df=None, pay_df=None) -> pd.DataFrame:
    """For a single month: who owed what, who paid, who didn't."""
    mem_df = members(conn) if mem_df is None else mem_df
    pay_df = payments(conn) if pay_df is None else pay_df
    pp = pay_df[pay_df["period"] == period].groupby("member_id")["amount"].sum().to_dict()
    rows = []
    for m in mem_df.to_dict("records"):
        due = float(m.get("monthly_due") or 0) if period in billable_periods(m, period) else 0.0
        paid = float(pp.get(m["id"], 0.0))
        if due == 0 and paid == 0:
            continue
        status = "Paid" if paid >= due and due > 0 else "Partial" if paid > 0 and paid < due else \
            "Advance / Extra" if due == 0 else "Unpaid"
        rows.append({"member_code": m["member_code"], "full_name": m["full_name"],
                     "university": m.get("university"), "district": m.get("district"),
                     "phone": m.get("phone"), "membership_type": m.get("membership_type"),
                     "due": due, "paid": paid, "balance": due - paid, "status": status})
    return pd.DataFrame(rows, columns=["member_code", "full_name", "university", "district", "phone",
                                       "membership_type", "due", "paid", "balance", "status"])


def collection_trend(conn, end: str, months: int = 12, mem_df=None, pay_df=None) -> pd.DataFrame:
    mem_df = members(conn) if mem_df is None else mem_df
    pay_df = payments(conn) if pay_df is None else pay_df
    rows = []
    for p in period_range(shift_period(end, -(months - 1)), end):
        exp = sum(float(m.get("monthly_due") or 0) for m in mem_df.to_dict("records")
                  if p in billable_periods(m, p))
        col = float(pay_df.loc[pay_df["period"] == p, "amount"].sum())
        rows.append({"Period": p, "Month": period_label(p), "Expected": exp, "Collected": col,
                     "Outstanding": max(exp - col, 0)})
    return pd.DataFrame(rows)


def unpaid_periods(member: dict, pay_df: pd.DataFrame, upto: str) -> list[str]:
    led = member_ledger(member, pay_df, upto)
    return led[(led["Due"] > 0) & (led["Paid"] < led["Due"])]["Period"].tolist()


# --------------------------------------------------------------------------------------
# Excel export
# --------------------------------------------------------------------------------------
def to_excel(sheets: dict[str, pd.DataFrame], title: str | None = None) -> bytes:
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    buf = io.BytesIO()
    head_fill = PatternFill("solid", fgColor="0D2B55")
    head_font = Font(bold=True, color="FFFFFF")
    title_font = Font(bold=True, size=14, color="0D2B55")
    thin = Side(style="thin", color="D0D5DD")
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        for name, df in sheets.items():
            start = 2 if title else 0
            df.to_excel(xw, sheet_name=name[:31], index=False, startrow=start)
            ws = xw.sheets[name[:31]]
            if title:
                ws.cell(row=1, column=1, value=f"{title} — {name}").font = title_font
            hr = start + 1
            for ci, col in enumerate(df.columns, 1):
                c = ws.cell(row=hr, column=ci)
                c.fill, c.font = head_fill, head_font
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                width = max([len(str(col))] + [len(str(v)) for v in df[col].head(500).tolist()]) + 2
                ws.column_dimensions[get_column_letter(ci)].width = min(max(width, 10), 50)
                is_money = pd.api.types.is_float_dtype(df[col])
                for r in range(hr + 1, hr + 1 + len(df)):
                    cell = ws.cell(row=r, column=ci)
                    cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
                    if is_money:
                        cell.number_format = "#,##0.00"
            ws.freeze_panes = ws.cell(row=hr + 1, column=1)
    return buf.getvalue()


# Everything above is the data layer; expose it as `db.<name>` for the UI code below.
db = types.SimpleNamespace(**{k: v for k, v in dict(globals()).items() if not k.startswith('__')})

# ======================================================================================
# USER INTERFACE
# ======================================================================================
import datetime as dt

import altair as alt
import pandas as pd
import streamlit as st


st.set_page_config(page_title="Member Database", page_icon="🎓", layout="wide")

NAVY, GOLD = "#0D2B55", "#C9A84C"
st.markdown(f"""
<style>
  .block-container {{padding-top: 1.6rem;}}
  [data-testid="stSidebar"] {{background: {NAVY};}}
  [data-testid="stSidebar"] * {{color: #fff !important;}}
  [data-testid="stMetric"] {{background:#fff;border:1px solid #e6e9ef;border-left:5px solid {GOLD};
      border-radius:10px;padding:12px 16px;box-shadow:0 1px 3px rgba(0,0,0,.05);}}
  [data-testid="stMetricLabel"] p {{color:#5b6475;font-weight:600;}}
  .uni-card {{background:#f5f7fb;border:1px solid #dfe5f0;border-left:5px solid {NAVY};
      border-radius:10px;padding:10px 16px;margin:4px 0 12px 0;font-size:.92rem;line-height:1.6}}
  .page-title {{color:{NAVY};font-weight:800;font-size:1.9rem;margin-bottom:0}}
  .page-sub {{color:#6b7280;margin-top:0;margin-bottom:1rem}}
</style>
""", unsafe_allow_html=True)


# --------------------------------------------------------------------------------------
# Setup & helpers
# --------------------------------------------------------------------------------------
@st.cache_resource
def get_conn():
    return db.connect()


conn = get_conn()
S = db.get_settings(conn)
CUR = S.get("currency", "¥")
TODAY = dt.date.today()
THIS_PERIOD = db.period_of(TODAY)


def money(x) -> str:
    return f"{CUR}{float(x or 0):,.2f}"


def money_col(label):
    return st.column_config.NumberColumn(label, format=f"{CUR}%.2f")


def title(text, sub=""):
    st.markdown(f"<p class='page-title'>{text}</p><p class='page-sub'>{sub}</p>", unsafe_allow_html=True)


def uni_label(row) -> str:
    cn = f" ({row['name_cn']})" if row.get("name_cn") else ""
    return f"{row['name_en']}{cn}"


def uni_card(u: dict):
    st.markdown(
        f"<div class='uni-card'><b>{u['name_en']}</b> &nbsp;{u.get('name_cn') or ''}<br>"
        f"🏛️ {u.get('ownership') or '—'} &nbsp;|&nbsp; 📍 {u.get('district') or '—'}<br>"
        f"🏠 {u.get('address') or '—'} &nbsp;|&nbsp; ✉️ Postal code {u.get('postal_code') or '—'}</div>",
        unsafe_allow_html=True)


def to_date(v):
    if v is None or (isinstance(v, float) and pd.isna(v)) or v == "":
        return None
    try:
        return dt.date.fromisoformat(str(v)[:10])
    except ValueError:
        return None


def period_options(back=36, forward=12):
    return db.period_range(db.shift_period(THIS_PERIOD, -back), db.shift_period(THIS_PERIOD, forward))


def member_picker(label, mem_df, key, help=None):
    if mem_df.empty:
        st.info("No members registered yet. Add one under **➕ Register Member**.")
        return None
    opts = mem_df["id"].tolist()
    lab = {r.id: f"{r.member_code} · {r.full_name} · {r.university or '—'}" for r in mem_df.itertuples()}
    return st.selectbox(label, opts, format_func=lambda i: lab[i], key=key, help=help)


def check_password() -> bool:
    try:
        pw = st.secrets.get("app_password")
    except Exception:
        pw = None
    if not pw or st.session_state.get("auth_ok"):
        return True
    title("🔒 Sign in", S.get("org_name", ""))
    entered = st.text_input("Password", type="password")
    if st.button("Sign in", type="primary"):
        if entered == pw:
            st.session_state["auth_ok"] = True
            st.rerun()
        st.error("Incorrect password.")
    return False


# --------------------------------------------------------------------------------------
# Member form (shared by Register and Edit)
# --------------------------------------------------------------------------------------
def member_form(prefix: str, d: dict | None = None, portal: bool = False) -> dict:
    d = d or {}
    unis = db.universities(conn)
    types = S["membership_types"] or {"Regular": 0}

    st.markdown("#### 🏫 University")
    c1, c2 = st.columns([1, 2])
    districts = ["All districts"] + sorted(unis["district"].dropna().unique().tolist())
    ownerships = ["All"] + sorted(unis["ownership"].dropna().unique().tolist())
    with c1:
        dist = st.selectbox("Filter by district", districts, key=f"{prefix}_dist")
        own = st.selectbox("Filter by ownership", ownerships, key=f"{prefix}_own")
    filt = unis.copy()
    if dist != "All districts":
        filt = filt[filt["district"] == dist]
    if own != "All":
        filt = filt[filt["ownership"] == own]
    cur_uid = d.get("university_id")
    if cur_uid and cur_uid not in filt["id"].tolist():
        filt = pd.concat([unis[unis["id"] == cur_uid], filt])
    uni_ids = filt["id"].tolist()
    ulabels = {r["id"]: uni_label(r) for r in unis.to_dict("records")}
    with c2:
        uid = st.selectbox(f"University * ({len(uni_ids)} listed)", uni_ids,
                           index=uni_ids.index(cur_uid) if cur_uid in uni_ids else None,
                           format_func=lambda i: ulabels.get(i, "?"),
                           placeholder="Select the member's university…", key=f"{prefix}_uni")
        if uid:
            uni_card(unis[unis["id"] == uid].iloc[0].to_dict())

    st.markdown("#### 👤 Personal details")
    c1, c2, c3 = st.columns(3)
    full_name = c1.text_input("Full name *", d.get("full_name") or "", key=f"{prefix}_name")
    g = d.get("gender")
    gender = c2.selectbox("Gender", db.GENDERS, index=db.GENDERS.index(g) if g in db.GENDERS else None,
                          placeholder="Select…", key=f"{prefix}_gender")
    dob = c3.date_input("Date of birth", to_date(d.get("date_of_birth")), min_value=dt.date(1940, 1, 1),
                        max_value=TODAY, format="YYYY-MM-DD", key=f"{prefix}_dob")
    c1, c2, c3 = st.columns(3)
    nationality = c1.text_input("Nationality", d.get("nationality") or "", key=f"{prefix}_nat")
    phone = c2.text_input("Phone (China)", d.get("phone") or "", placeholder="+86 …", key=f"{prefix}_phone")
    email = c3.text_input("Email", d.get("email") or "", key=f"{prefix}_email")
    c1, c2, c3 = st.columns(3)
    wechat = c1.text_input("WeChat ID", d.get("wechat_id") or "", key=f"{prefix}_wechat")
    em_name = c2.text_input("Emergency contact name", d.get("emergency_name") or "", key=f"{prefix}_emn")
    em_phone = c3.text_input("Emergency contact phone", d.get("emergency_phone") or "", key=f"{prefix}_emp")

    st.markdown("#### 📚 Academic details")
    c1, c2, c3 = st.columns(3)
    student_id = c1.text_input("Student ID", d.get("student_id") or "", key=f"{prefix}_sid")
    dl = d.get("degree_level")
    degree = c2.selectbox("Level of study", db.DEGREE_LEVELS,
                          index=db.DEGREE_LEVELS.index(dl) if dl in db.DEGREE_LEVELS else None,
                          placeholder="Select…", key=f"{prefix}_deg")
    field = c3.text_input("Programme / field of study", d.get("field_of_study") or "", key=f"{prefix}_field")
    c1, c2, c3, c4 = st.columns(4)
    years = list(range(TODAY.year - 12, TODAY.year + 9))
    ey, gy = d.get("enrollment_year"), d.get("expected_graduation")
    ey = int(ey) if ey and not pd.isna(ey) else None
    gy = int(gy) if gy and not pd.isna(gy) else None
    enroll = c1.selectbox("Enrollment year", years, index=years.index(ey) if ey in years else None,
                          placeholder="Year…", key=f"{prefix}_ey")
    grad = c2.selectbox("Expected graduation", years, index=years.index(gy) if gy in years else None,
                        placeholder="Year…", key=f"{prefix}_gy")
    sc = d.get("scholarship")
    scholarship = c3.selectbox("Scholarship / funding", db.SCHOLARSHIPS,
                               index=db.SCHOLARSHIPS.index(sc) if sc in db.SCHOLARSHIPS else None,
                               placeholder="Select…", key=f"{prefix}_sch")
    rs = d.get("residence")
    residence = c4.selectbox("Residence", db.RESIDENCE,
                             index=db.RESIDENCE.index(rs) if rs in db.RESIDENCE else None,
                             placeholder="Select…", key=f"{prefix}_res")

    st.markdown("#### 💳 Membership & dues")
    c1, c2, c3, c4 = st.columns(4)
    tnames = list(types.keys())
    mt = d.get("membership_type")
    mtype = c1.selectbox("Membership type", tnames, index=tnames.index(mt) if mt in tnames else 0,
                         key=f"{prefix}_mtype")
    due_default = float(d["monthly_due"]) if d.get("monthly_due") is not None and mt == mtype \
        else float(types.get(mtype, 0))
    # key includes the type so the amount resets to that type's default when the type changes
    if portal:  # members can't set their own dues, dates or status
        monthly_due, join, status = due_default, TODAY, db.PENDING
        c2.metric("Monthly due", money(due_default))
        c3.caption("Your registration will be reviewed by an administrator. "
                   "Dues start from the month you are approved.")
    else:
        monthly_due = c2.number_input(f"Monthly due ({CUR})", min_value=0.0, step=5.0, value=due_default,
                                      key=f"{prefix}_due_{mtype}")
        join = c3.date_input("Join date (billing starts) *", to_date(d.get("join_date")) or TODAY,
                             format="YYYY-MM-DD", key=f"{prefix}_join")
        stt = d.get("status") or "Active"
        status = c4.selectbox("Status", db.STATUSES, index=db.STATUSES.index(stt), key=f"{prefix}_status")
    exit_date = None
    if status in db.EXIT_STATUSES:
        exit_date = st.date_input("Date left / graduated (billing stops after this month) *",
                                  to_date(d.get("exit_date")) or TODAY, format="YYYY-MM-DD",
                                  key=f"{prefix}_exit")
    notes = st.text_area("Anything else we should know?" if portal else "Notes", d.get("notes") or "",
                         key=f"{prefix}_notes", height=80)

    return dict(full_name=full_name, gender=gender, date_of_birth=dob, nationality=nationality, phone=phone,
                email=email, wechat_id=wechat, university_id=uid, student_id=student_id, degree_level=degree,
                field_of_study=field, enrollment_year=enroll, expected_graduation=grad, scholarship=scholarship,
                residence=residence, membership_type=mtype, monthly_due=monthly_due, join_date=join,
                status=status, exit_date=exit_date, emergency_name=em_name, emergency_phone=em_phone,
                notes=notes)


# --------------------------------------------------------------------------------------
# Pages
# --------------------------------------------------------------------------------------
def page_dashboard():
    title("📊 Dashboard", f"{S.get('org_name')} · as of {TODAY:%d %B %Y}")
    mem, pay = db.members(conn), db.payments(conn)
    if mem.empty:
        st.info("Welcome! Start by registering members under **➕ Register Member**. "
                f"{len(db.universities(conn))} universities are loaded and ready in the dropdown.")
        return
    npend = int((mem["status"] == db.PENDING).sum())
    if npend:
        st.warning(f"📝 {npend} self-registration(s) from the member page are waiting for approval "
                   "— open **👥 Members** to review them.")
    summ = db.dues_summary(conn, THIS_PERIOD, mem, pay)
    month = db.monthly_collection(conn, THIS_PERIOD, mem, pay)
    c = st.columns(5)
    c[0].metric("Total members", len(mem))
    c[1].metric("Active members", int((mem["status"] == "Active").sum()))
    c[2].metric(f"Due · {db.period_label(THIS_PERIOD)}", money(month["due"].sum()))
    c[3].metric(f"Collected · {db.period_label(THIS_PERIOD)}", money(month["paid"].sum()))
    out = summ["balance"].clip(lower=0).sum()
    c[4].metric("Total outstanding", money(out),
                f"{int((summ['balance'] > 0).sum())} members owing", delta_color="inverse")

    left, right = st.columns(2)
    with left:
        st.markdown("##### Collections – last 12 months")
        tr = db.collection_trend(conn, THIS_PERIOD, 12, mem, pay)
        long = tr.melt(id_vars=["Period", "Month"], value_vars=["Expected", "Collected"],
                       var_name="Type", value_name="Amount")
        chart = alt.Chart(long).mark_bar().encode(
            x=alt.X("Month:N", sort=tr["Month"].tolist(), title=None),
            xOffset="Type:N",
            y=alt.Y("Amount:Q", title=CUR),
            color=alt.Color("Type:N", scale=alt.Scale(range=[NAVY, GOLD]), legend=alt.Legend(orient="top")),
            tooltip=["Month", "Type", alt.Tooltip("Amount:Q", format=",.2f")]).properties(height=320)
        st.altair_chart(chart, use_container_width=True)
    with right:
        st.markdown("##### Members by university")
        bu = mem.groupby("university").size().reset_index(name="Members").sort_values("Members", ascending=False)
        chart = alt.Chart(bu).mark_bar(color=NAVY).encode(
            x=alt.X("Members:Q"), y=alt.Y("university:N", sort="-x", title=None),
            tooltip=["university", "Members"]).properties(height=320)
        st.altair_chart(chart, use_container_width=True)

    left, right = st.columns(2)
    with left:
        st.markdown("##### Members by district")
        bd = mem.groupby("district").size().reset_index(name="Members")
        st.altair_chart(alt.Chart(bd).mark_arc(innerRadius=60).encode(
            theta="Members:Q", color=alt.Color("district:N", title="District"),
            tooltip=["district", "Members"]).properties(height=300), use_container_width=True)
    with right:
        st.markdown("##### Top outstanding balances")
        top = summ[summ["balance"] > 0].sort_values("balance", ascending=False).head(10)
        if top.empty:
            st.success("🎉 Everyone is up to date!")
        else:
            st.dataframe(top[["member_code", "full_name", "university", "months_in_arrears", "balance"]],
                         hide_index=True, use_container_width=True,
                         column_config={"balance": money_col("Balance"),
                                        "months_in_arrears": "Months owed",
                                        "full_name": "Name", "member_code": "Code"})


def page_register():
    title("➕ Register Member", "Fill in the member's details. Fields marked * are required.")
    ver = st.session_state.setdefault("reg_ver", 0)
    if st.session_state.get("reg_done"):
        st.success(st.session_state.pop("reg_done"))
    data = member_form(f"reg{ver}")
    st.divider()
    if st.button("💾 Save member", type="primary", use_container_width=True):
        errs = db.validate_member(conn, data)
        if errs:
            for e in errs:
                st.error(e)
        else:
            code = db.add_member(conn, data)
            st.session_state["reg_done"] = f"✅ {data['full_name']} registered as **{code}**."
            st.session_state["reg_ver"] = ver + 1  # fresh, empty form
            st.rerun()


def page_members():
    title("👥 Members Directory", "Search, filter, view and edit member records.")
    mem = db.members(conn)
    if mem.empty:
        st.info("No members yet.")
        return
    pend = mem[mem["status"] == db.PENDING]
    if not pend.empty:
        with st.container(border=True):
            st.markdown(f"#### 📝 Waiting for approval ({len(pend)})")
            st.caption("These people registered themselves through the QR code / member page. "
                       "Dues start from the join date you approve them with.")
            for r in pend.to_dict("records"):
                a, b, c, d_ = st.columns([3, 1.4, 1, 1])
                a.markdown(f"**{r['full_name']}** · `{r['member_code']}`  \n"
                           f"{r.get('university') or '—'} · {r.get('phone') or 'no phone'} · "
                           f"{r.get('membership_type')} ({money(r.get('monthly_due'))}/month)")
                jd = b.date_input("Join date", TODAY, key=f"apj{r['id']}", format="YYYY-MM-DD",
                                  label_visibility="collapsed")
                if c.button("✅ Approve", key=f"ap{r['id']}", use_container_width=True):
                    db.approve_member(conn, r["id"], jd)
                    st.toast(f"{r['full_name']} approved")
                    st.rerun()
                if d_.button("✖ Reject", key=f"rj{r['id']}", use_container_width=True):
                    db.delete_member(conn, r["id"])
                    st.toast(f"{r['full_name']} removed")
                    st.rerun()
    summ = db.dues_summary(conn, THIS_PERIOD, mem, db.payments(conn)).set_index("id")
    mem = mem.join(summ[["balance", "months_in_arrears"]], on="id")

    c1, c2, c3, c4, c5 = st.columns([2, 2, 1.4, 1.2, 1.2])
    q = c1.text_input("🔎 Search", placeholder="Name, code, phone, email, student ID…")
    unis = c2.multiselect("University", sorted(mem["university"].dropna().unique()))
    dists = c3.multiselect("District", sorted(mem["district"].dropna().unique()))
    stats = c4.multiselect("Status", db.STATUSES)
    owing = c5.selectbox("Dues", ["All", "Owing", "Up to date"])
    f = mem
    if q:
        ql = q.lower()
        hay = f[["full_name", "member_code", "phone", "email", "student_id", "wechat_id"]].fillna("").astype(str)
        f = f[hay.apply(lambda r: ql in " ".join(r).lower(), axis=1)]
    if unis:
        f = f[f["university"].isin(unis)]
    if dists:
        f = f[f["district"].isin(dists)]
    if stats:
        f = f[f["status"].isin(stats)]
    if owing == "Owing":
        f = f[f["balance"] > 0]
    elif owing == "Up to date":
        f = f[f["balance"] <= 0]

    show = ["member_code", "full_name", "gender", "university", "district", "degree_level", "phone", "email",
            "membership_type", "monthly_due", "status", "join_date", "balance", "months_in_arrears"]
    st.caption(f"Showing {len(f)} of {len(mem)} members")
    st.dataframe(f[show], hide_index=True, use_container_width=True, height=380,
                 column_config={"member_code": "Code", "full_name": "Name", "degree_level": "Level",
                                "membership_type": "Type", "monthly_due": money_col("Monthly due"),
                                "balance": money_col("Balance"), "months_in_arrears": "Months owed",
                                "join_date": "Joined"})
    st.download_button("⬇️ Download this list (Excel)",
                       db.to_excel({"Members": f[show + ["wechat_id", "student_id", "nationality",
                                                         "field_of_study", "scholarship"]]},
                                   S.get("org_name")),
                       f"members_{TODAY}.xlsx")

    st.divider()
    st.subheader("✏️ View / edit a member")
    mid = member_picker("Select member", f if not f.empty else mem, "edit_pick")
    if not mid:
        return
    m = db.get_member(conn, mid)
    with st.expander(f"Edit {m['member_code']} · {m['full_name']}", expanded=False):
        data = member_form(f"edit{mid}", m)
        b1, b2 = st.columns([1, 1])
        if b1.button("💾 Save changes", type="primary", use_container_width=True):
            errs = db.validate_member(conn, data, mid)
            if errs:
                for e in errs:
                    st.error(e)
            else:
                db.update_member(conn, mid, data)
                st.toast("Member updated ✅")
                st.rerun()
        with b2:
            sure = st.checkbox("I understand this deletes the member and all their payments")
            if st.button("🗑️ Delete member", disabled=not sure, use_container_width=True):
                db.delete_member(conn, mid)
                st.toast("Member deleted")
                st.rerun()


def page_payments():
    title("💰 Record Payment", "Record monthly dues. Unpaid months are pre-selected, oldest first.")
    mem, pay = db.members(conn), db.payments(conn)
    mid = member_picker("Member", mem, "pay_pick")
    if not mid:
        return
    m = db.get_member(conn, mid)
    due = float(m.get("monthly_due") or 0)
    unpaid = db.unpaid_periods(m, pay, THIS_PERIOD)
    led = db.member_ledger(m, pay, THIS_PERIOD)
    owed = float(led.loc[led["Period"] <= THIS_PERIOD, "Due"].sum()) - float(
        pay.loc[pay["member_id"] == mid, "amount"].sum())

    c = st.columns(4)
    c[0].metric("Monthly due", money(due))
    c[1].metric("Months owed", len(unpaid))
    c[2].metric("Balance to date", money(owed))
    c[3].metric("Status", m.get("status") or "—")

    with st.form("pay_form", clear_on_submit=True):
        opts = sorted(set(period_options()) | set(unpaid))
        months = st.multiselect("Month(s) this payment covers", opts, default=unpaid[:1] if unpaid else [],
                                format_func=db.period_label,
                                help="Pick several months to record a multi-month or advance payment.")
        c1, c2, c3 = st.columns(3)
        amt = c1.number_input(f"Amount per month ({CUR})", min_value=0.0, value=due, step=5.0,
                              help="Enter less than the due amount to record a partial payment.")
        pdate = c2.date_input("Payment date", TODAY, format="YYYY-MM-DD")
        method = c3.selectbox("Method", db.PAY_METHODS)
        c1, c2 = st.columns(2)
        ref = c1.text_input("Receipt / transaction reference")
        by = c2.text_input("Received by")
        notes = st.text_input("Notes")
        submitted = st.form_submit_button("💾 Record payment", type="primary", use_container_width=True)
    if submitted:
        if not months:
            st.error("Select at least one month.")
        elif amt <= 0:
            st.error("Amount must be greater than zero.")
        else:
            n = db.add_payments(conn, mid, months, amt, pdate, method, ref, by, notes)
            st.toast(f"Recorded {n} month(s) · total {money(n * amt)} ✅")
            st.rerun()
    if unpaid:
        st.caption(f"Tip: pay all {len(unpaid)} owed months at once → total {money(due * len(unpaid))}. "
                   f"Owed: {', '.join(db.period_label(p) for p in unpaid)}")

    st.markdown("##### Payment history")
    hist = pay[pay["member_id"] == mid]
    if hist.empty:
        st.caption("No payments yet.")
    else:
        h = hist[["id", "period", "amount", "payment_date", "method", "reference", "received_by", "notes"]].copy()
        h["period"] = h["period"].map(db.period_label)
        st.dataframe(h.drop(columns="id"), hide_index=True, use_container_width=True,
                     column_config={"amount": money_col("Amount"), "period": "For month"})
        with st.expander("Remove a payment entered by mistake"):
            pid = st.selectbox("Payment", hist["id"].tolist(),
                               format_func=lambda i: (lambda r: f"#{i} · {db.period_label(r.period)} · "
                                                      f"{money(r.amount)} · {r.payment_date}")(
                                   hist[hist["id"] == i].iloc[0]))
            if st.button("🗑️ Delete payment"):
                db.delete_payment(conn, pid)
                st.toast("Payment removed")
                st.rerun()


def page_reports():
    title("📑 Reports", "Monthly dues, arrears, member statements, university summaries and exports.")
    mem, pay = db.members(conn), db.payments(conn)
    if mem.empty:
        st.info("No members yet.")
        return
    tabs = st.tabs(["📅 Monthly dues", "⚠️ Arrears", "🧾 Member statement", "🏫 By university",
                    "📒 Payment register", "📦 Full export"])
    org = S.get("org_name")

    # ---- Monthly dues -------------------------------------------------------------
    with tabs[0]:
        opts = period_options(36, 12)
        p = st.selectbox("Month", opts, index=opts.index(THIS_PERIOD), format_func=db.period_label, key="rm_p")
        mc = db.monthly_collection(conn, p, mem, pay)
        c = st.columns(5)
        c[0].metric("Members billed", int((mc["due"] > 0).sum()))
        c[1].metric("Total due", money(mc["due"].sum()))
        c[2].metric("Collected", money(mc["paid"].sum()))
        c[3].metric("Outstanding", money(mc["balance"].clip(lower=0).sum()))
        rate = mc["paid"].clip(upper=mc["due"]).sum() / mc["due"].sum() * 100 if mc["due"].sum() else 0
        c[4].metric("Collection rate", f"{rate:.0f}%")
        sf = st.multiselect("Show status", ["Paid", "Partial", "Unpaid", "Advance / Extra"],
                            default=["Paid", "Partial", "Unpaid", "Advance / Extra"], key="rm_s")
        mc = mc[mc["status"].isin(sf)]
        st.dataframe(mc.style.map(lambda v: "color:#15803d;font-weight:600" if v == "Paid" else
                                  "color:#b45309;font-weight:600" if v == "Partial" else
                                  "color:#b91c1c;font-weight:600" if v == "Unpaid" else "", subset=["status"]),
                     hide_index=True, use_container_width=True,
                     column_config={"due": money_col("Due"), "paid": money_col("Paid"),
                                    "balance": money_col("Balance"), "full_name": "Name",
                                    "member_code": "Code", "membership_type": "Type"})
        st.download_button("⬇️ Download monthly dues report", db.to_excel(
            {f"Dues {db.period_label(p)}": mc}, org), f"monthly_dues_{p}.xlsx", key="rm_dl")

    # ---- Arrears ------------------------------------------------------------------
    with tabs[1]:
        opts = period_options(36, 0)
        p = st.selectbox("Balances as of end of", opts, index=opts.index(THIS_PERIOD),
                         format_func=db.period_label, key="ra_p")
        sm = db.dues_summary(conn, p, mem, pay)
        ar = sm[sm["balance"] > 0].sort_values(["months_in_arrears", "balance"], ascending=False)
        mx = int(ar["months_in_arrears"].max()) if not ar.empty else 1
        minm = st.slider("Minimum months owed", 1, mx, 1) if mx > 1 else 1
        ar = ar[ar["months_in_arrears"] >= minm]
        c = st.columns(3)
        c[0].metric("Members in arrears", len(ar))
        c[1].metric("Total arrears", money(ar["balance"].sum()))
        c[2].metric("Average per member", money(ar["balance"].mean() if len(ar) else 0))
        cols = ["member_code", "full_name", "university", "phone", "email", "status", "monthly_due",
                "months_in_arrears", "oldest_unpaid", "expected", "paid", "balance", "last_payment"]
        st.dataframe(ar[cols], hide_index=True, use_container_width=True,
                     column_config={"monthly_due": money_col("Monthly due"), "expected": money_col("Expected"),
                                    "paid": money_col("Paid"), "balance": money_col("Balance"),
                                    "months_in_arrears": "Months owed", "oldest_unpaid": "Oldest unpaid",
                                    "full_name": "Name", "member_code": "Code"})
        st.download_button("⬇️ Download arrears report", db.to_excel({"Arrears": ar[cols]}, org),
                           f"arrears_{p}.xlsx", key="ra_dl")

    # ---- Member statement ---------------------------------------------------------
    with tabs[2]:
        mid = member_picker("Member", mem, "rs_m")
        if mid:
            m = db.get_member(conn, mid)
            led = db.member_ledger(m, pay, THIS_PERIOD)
            exp = float(led.loc[led["Period"] <= THIS_PERIOD, "Due"].sum())
            paid = float(led["Paid"].sum())
            a, b = st.columns([1.3, 1])
            with a:
                st.markdown(f"### {m['full_name']}  \n`{m['member_code']}` · {m.get('status')}")
                st.markdown(
                    f"**University:** {m.get('university') or '—'} {m.get('university_cn') or ''}  \n"
                    f"**District:** {m.get('district') or '—'}  \n"
                    f"**Level / programme:** {m.get('degree_level') or '—'} · {m.get('field_of_study') or '—'}  \n"
                    f"**Phone:** {m.get('phone') or '—'} · **WeChat:** {m.get('wechat_id') or '—'} · "
                    f"**Email:** {m.get('email') or '—'}  \n"
                    f"**Membership:** {m.get('membership_type')} · {money(m.get('monthly_due'))}/month · "
                    f"joined {m.get('join_date')}")
            with b:
                st.metric("Total due to date", money(exp))
                st.metric("Total paid", money(paid))
                st.metric("Balance", money(exp - paid), delta_color="inverse")
            st.dataframe(led.drop(columns="Period").style.map(
                lambda v: "color:#15803d" if v == "Paid" else "color:#b45309" if v == "Partial" else
                "color:#b91c1c" if v == "Unpaid" else "", subset=["Status"]),
                hide_index=True, use_container_width=True,
                column_config={k: money_col(k) for k in ["Due", "Paid", "Balance", "Running Balance"]})
            profile = pd.DataFrame([{k: m.get(k) for k in [
                "member_code", "full_name", "gender", "nationality", "phone", "email", "wechat_id",
                "university", "university_cn", "district", "student_id", "degree_level", "field_of_study",
                "membership_type", "monthly_due", "join_date", "status"]}]).T.reset_index()
            profile.columns = ["Field", "Value"]
            st.download_button("⬇️ Download statement", db.to_excel(
                {"Profile": profile, "Statement": led.drop(columns="Period")}, f"{org} · {m['full_name']}"),
                f"statement_{m['member_code']}_{TODAY}.xlsx", key="rs_dl")

    # ---- By university ------------------------------------------------------------
    with tabs[3]:
        sm = db.dues_summary(conn, THIS_PERIOD, mem, pay)
        g = sm.groupby(["university", "district"], dropna=False).agg(
            members=("id", "count"), active=("status", lambda s: (s == "Active").sum()),
            expected=("expected", "sum"), paid=("paid", "sum"), outstanding=("balance", "sum")
        ).reset_index().sort_values("members", ascending=False)
        g["collection_rate_%"] = (g["paid"] / g["expected"].where(g["expected"] > 0) * 100).round(1)
        st.dataframe(g, hide_index=True, use_container_width=True,
                     column_config={"expected": money_col("Expected"), "paid": money_col("Paid"),
                                    "outstanding": money_col("Outstanding")})
        st.altair_chart(alt.Chart(g).mark_bar(color=GOLD).encode(
            x=alt.X("outstanding:Q", title=f"Outstanding ({CUR})"), y=alt.Y("university:N", sort="-x", title=None),
            tooltip=["university", "members", alt.Tooltip("outstanding:Q", format=",.2f")]),
            use_container_width=True)
        st.download_button("⬇️ Download university summary", db.to_excel({"By University": g}, org),
                           f"university_summary_{TODAY}.xlsx", key="ru_dl")

    # ---- Payment register ---------------------------------------------------------
    with tabs[4]:
        c1, c2, c3 = st.columns(3)
        d1 = c1.date_input("From", TODAY.replace(day=1), format="YYYY-MM-DD", key="rp_d1")
        d2 = c2.date_input("To", TODAY, format="YYYY-MM-DD", key="rp_d2")
        meth = c3.multiselect("Method", db.PAY_METHODS, key="rp_m")
        r = pay[(pay["payment_date"] >= str(d1)) & (pay["payment_date"] <= str(d2))]
        if meth:
            r = r[r["method"].isin(meth)]
        r = r[["payment_date", "member_code", "full_name", "university", "period", "amount", "method",
               "reference", "received_by", "notes"]].copy()
        r["period"] = r["period"].map(db.period_label)
        c = st.columns(3)
        c[0].metric("Payments", len(r))
        c[1].metric("Total received", money(r["amount"].sum()))
        c[2].metric("Members paying", r["member_code"].nunique())
        st.dataframe(r, hide_index=True, use_container_width=True,
                     column_config={"amount": money_col("Amount"), "period": "For month", "full_name": "Name"})
        if not r.empty:
            bym = r.groupby("method")["amount"].sum().reset_index()
            st.altair_chart(alt.Chart(bym).mark_bar(color=NAVY).encode(
                x=alt.X("method:N", title=None), y=alt.Y("amount:Q", title=CUR)).properties(height=220),
                use_container_width=True)
        st.download_button("⬇️ Download payment register", db.to_excel({"Payments": r}, org),
                           f"payments_{d1}_{d2}.xlsx", key="rp_dl")

    # ---- Full export --------------------------------------------------------------
    with tabs[5]:
        st.write("One Excel workbook with every member, every payment, the current dues summary, "
                 "this month's dues list and the university list.")
        sm = db.dues_summary(conn, THIS_PERIOD, mem, pay)
        book = {
            "Members": mem.drop(columns=["university_id"]),
            "Dues Summary": sm.drop(columns=["id"]),
            f"Dues {db.period_label(THIS_PERIOD)}": db.monthly_collection(conn, THIS_PERIOD, mem, pay),
            "Payments": pay.drop(columns=["member_id"]),
            "Universities": db.universities(conn).drop(columns=["id"]),
        }
        st.download_button("⬇️ Download full workbook", db.to_excel(book, org),
                           f"member_database_{TODAY}.xlsx", type="primary", key="rf_dl")
        st.download_button("⬇️ Members (CSV)", mem.to_csv(index=False).encode("utf-8-sig"),
                           f"members_{TODAY}.csv", key="rf_csv")


def page_settings():
    title("⚙️ Settings", "Organisation details, membership types, universities and backups.")
    t1, tq, t2, t3 = st.tabs(["General & dues", "📱 QR code", "Universities", "Backup & restore"])
    with tq:
        page_qr()
    with t1:
        with st.form("settings"):
            org = st.text_input("Organisation name", S.get("org_name"))
            cur = st.text_input("Currency symbol", S.get("currency"), max_chars=4)
            st.markdown("**Membership types & default monthly dues** — edit, add or remove rows.")
            mt = pd.DataFrame([{"Membership type": k, "Default monthly due": float(v)}
                               for k, v in S["membership_types"].items()],
                              columns=["Membership type", "Default monthly due"])
            mt = st.data_editor(mt, num_rows="dynamic", use_container_width=True, hide_index=True)
            if st.form_submit_button("💾 Save settings", type="primary"):
                types = {str(r["Membership type"]).strip(): float(r["Default monthly due"] or 0)
                         for _, r in mt.dropna(subset=["Membership type"]).iterrows()
                         if str(r["Membership type"]).strip()}
                if not types:
                    st.error("Keep at least one membership type.")
                else:
                    db.save_setting(conn, "org_name", org)
                    db.save_setting(conn, "currency", cur or "¥")
                    db.save_setting(conn, "membership_types", types)
                    st.toast("Settings saved ✅")
                    st.rerun()
        st.caption("Changing a default does not change existing members' dues — edit a member to change theirs.")

    with t2:
        unis = db.universities(conn)
        st.caption(f"{len(unis)} universities available in the registration dropdown.")
        st.dataframe(unis.drop(columns="id"), hide_index=True, use_container_width=True, height=300)
        with st.expander("➕ Add a university"):
            with st.form("add_uni", clear_on_submit=True):
                c1, c2 = st.columns(2)
                en = c1.text_input("Name (English) *")
                cn = c2.text_input("Name (Chinese)")
                own = c1.selectbox("Ownership", ["Public", "Private"])
                dists = sorted(unis["district"].dropna().unique().tolist())
                dist = c2.selectbox("District", dists + ["Other…"])
                other = c2.text_input("If other district, type it")
                addr = c1.text_input("Address")
                post = c2.text_input("Postal code")
                if st.form_submit_button("Add university"):
                    if not en.strip():
                        st.error("Name is required.")
                    else:
                        try:
                            db.add_university(conn, en, cn, own, other if dist == "Other…" else dist, addr, post)
                            st.toast("University added ✅")
                            st.rerun()
                        except Exception as e:
                            st.error(f"Could not add: {e}")
        with st.expander("🔄 Update list from Excel"):
            up = st.file_uploader("Excel file with the same columns as the original list", type=["xlsx"])
            if up and st.button("Import"):
                try:
                    n = db.import_universities(conn, up)
                    st.success(f"Imported / updated {n} universities.")
                except Exception as e:
                    st.error(str(e))

    with t3:
        st.warning("On Streamlit Community Cloud the app's storage is reset whenever the app restarts or "
                   "redeploys. **Download a backup regularly** and restore it after a restart.")
        st.download_button("⬇️ Download database backup (.db)", db.DB_PATH.read_bytes(),
                           f"members_backup_{TODAY}.db", type="primary")
        up = st.file_uploader("Restore from a backup (.db)", type=["db", "sqlite"])
        if up and st.button("♻️ Restore (replaces all current data)"):
            import sqlite3, tempfile
            raw = up.getvalue()
            with tempfile.NamedTemporaryFile(suffix=".db", delete=False) as tmp:
                tmp.write(raw)
            try:
                t = sqlite3.connect(tmp.name)
                t.execute("SELECT COUNT(*) FROM members").fetchone()
                t.close()
            except Exception:
                st.error("That file is not a valid backup from this app.")
            else:
                get_conn().close()
                db.DB_PATH.write_bytes(raw)
                st.cache_resource.clear()
                st.success("Backup restored.")
                st.rerun()


def has_password() -> bool:
    try:
        return bool(st.secrets.get("app_password"))
    except Exception:
        return False


def page_qr():
    st.markdown("Members scan this code with their phone camera or WeChat to open the **member page**, "
                "where they can **register themselves** and **check their own dues**. "
                "They cannot see other members or the admin screens.")
    url = st.text_input("Your app's web address", S.get("app_url", ""),
                        placeholder="e.g. https://your-app-name.streamlit.app",
                        help="Copy it from the browser address bar after deploying the app.")
    if url != S.get("app_url", ""):
        db.save_setting(conn, "app_url", url.strip())
    if not url.strip():
        st.info("Deploy the app (for example on Streamlit Community Cloud), paste its address above, "
                "and the QR code appears here.")
        return
    link = db.member_link(url)
    if not has_password():
        st.error("⚠️ No admin password is set. Anyone who removes `?page=member` from the link could open "
                 "the admin screens. Add `app_password` in your app's Secrets before sharing the QR code.")
    try:
        qr = db.qr_image(link)
    except ImportError:
        st.error("The `qrcode` package is missing — add `qrcode[pil]` to requirements.txt.")
        return
    a, b = st.columns([1, 1.3])
    with a:
        st.image(db.png_bytes(qr), caption="Member page QR code", width=300)
    with b:
        st.markdown("**Link inside the code**")
        st.code(link, language=None)
        st.link_button("Open the member page", link)
        st.download_button("⬇️ Download QR code (PNG)", db.png_bytes(qr), "member_qr_code.png",
                           mime="image/png", use_container_width=True)
        st.download_button("🖨️ Download printable poster (PNG)",
                           db.qr_poster(qr, S.get("org_name", ""), link), "member_qr_poster.png",
                           mime="image/png", type="primary", use_container_width=True)
        st.caption("Share the PNG in your WeChat group, or print the poster for meetings and noticeboards.")


def page_portal():
    st.markdown(f"<p class='page-title'>🎓 {S.get('org_name')}</p>"
                "<p class='page-sub'>Member page — register as a member or check your dues.</p>",
                unsafe_allow_html=True)
    t1, t2 = st.tabs(["📝 Register", "💳 Check my dues"])
    with t1:
        ver = st.session_state.setdefault("portal_ver", 0)
        if st.session_state.get("portal_done"):
            st.success(st.session_state.pop("portal_done"))
            st.balloons()
        data = member_form(f"portal{ver}", portal=True)
        st.divider()
        if st.button("📨 Submit registration", type="primary", use_container_width=True):
            errs = db.validate_member(conn, data)
            if not (data.get("phone") or "").strip():
                errs.append("Please enter your phone number — you'll need it to check your dues.")
            if errs:
                for e in errs:
                    st.error(e)
            else:
                code = db.add_member(conn, data)
                st.session_state["portal_done"] = (
                    f"Thank you, {data['full_name']}! Your registration was received. "
                    f"Your member code is **{code}** — please save it. You'll use it with the last "
                    "4 digits of your phone number to check your dues.")
                st.session_state["portal_ver"] = ver + 1
                st.rerun()
    with t2:
        st.caption("Enter your member code and the last 4 digits of the phone number you registered with.")
        with st.form("portal_lookup"):
            c1, c2 = st.columns(2)
            code = c1.text_input("Member code", placeholder="MBR-0001")
            last4 = c2.text_input("Last 4 digits of your phone", max_chars=4, type="password")
            go = st.form_submit_button("Check my dues", type="primary", use_container_width=True)
        if go:
            tries = st.session_state.get("portal_tries", 0)
            if tries >= 5:
                st.error("Too many attempts. Please refresh the page and try again later.")
                return
            m = db.portal_lookup(conn, code, last4)
            if not m:
                st.session_state["portal_tries"] = tries + 1
                st.error("No match. Check your member code and phone digits, or contact an administrator.")
                return
            st.session_state["portal_tries"] = 0
            st.markdown(f"### Hello, {m['full_name']} 👋")
            if m.get("status") == db.PENDING:
                st.info("Your registration is waiting for approval. Dues start once you are approved.")
                return
            led = db.member_ledger(m, db.payments(conn), THIS_PERIOD)
            exp = float(led.loc[led["Period"] <= THIS_PERIOD, "Due"].sum())
            paid = float(led["Paid"].sum())
            owed = led[(led["Due"] > 0) & (led["Paid"] < led["Due"]) & (led["Period"] <= THIS_PERIOD)]
            c = st.columns(3)
            c[0].metric("Monthly due", money(m.get("monthly_due")))
            c[1].metric("Months owed", len(owed))
            bal = exp - paid
            c[2].metric("Balance", money(bal) if bal > 0 else
                        ("Up to date ✅" if bal == 0 else f"{money(-bal)} credit"))
            if len(owed):
                st.warning("Unpaid: " + ", ".join(owed["Month"]))
            st.dataframe(led.drop(columns=["Period", "Running Balance"]), hide_index=True,
                         use_container_width=True,
                         column_config={k: money_col(k) for k in ["Due", "Paid", "Balance"]})


# --------------------------------------------------------------------------------------
# Navigation
# --------------------------------------------------------------------------------------
PAGES = {
    "📊 Dashboard": page_dashboard,
    "➕ Register Member": page_register,
    "👥 Members": page_members,
    "💰 Record Payment": page_payments,
    "📑 Reports": page_reports,
    "⚙️ Settings": page_settings,
}

if st.query_params.get("page") == "member":
    st.markdown("<style>[data-testid='stSidebar'],[data-testid='stSidebarCollapsedControl']"
                "{display:none}</style>", unsafe_allow_html=True)
    page_portal()
elif check_password():
    with st.sidebar:
        st.markdown(f"### 🎓 {S.get('org_name')}")
        st.caption("Member Database & Dues")
        page = st.radio("Navigate", list(PAGES), label_visibility="collapsed")
        st.divider()
        n = conn.execute("SELECT COUNT(*) FROM members").fetchone()[0]
        st.caption(f"{n} members · {TODAY:%d %b %Y}")
    PAGES[page]()
